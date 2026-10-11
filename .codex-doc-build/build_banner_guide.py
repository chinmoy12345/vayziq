from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"H:\NEXTJS\vayziq\docs\Vayziq_Banner_Source_Guide.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.12

for name, size in [("Title", 25), ("Heading 1", 16), ("Heading 2", 12)]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.keep_with_next = True

styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_margin(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        tblBorders.append(tag)
    tblPr.append(tblBorders)

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].width = Inches(widths[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(hdr[i], "252525")
        set_cell_margin(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(8.5)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trPr.append(repeat)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cells[i])
            if r_idx % 2:
                shade(cells[i], "F7F7F7")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Vayziq Banner Source Guide")

p = doc.add_paragraph()
p.add_run("Purpose  ").bold = True
p.add_run("This guide shows where each storefront banner comes from, which Admin placement controls it, and which code fallback appears when no active database banner exists.")

p = doc.add_paragraph()
p.add_run("Main conclusion  ").bold = True
p.add_run("Home, Offer Zone, Shop, Search and category collection banners can be managed from Admin Banners. About, Contact, legal pages and the Offers listing hero are currently code based and require a code or asset change.")

doc.add_heading("Banner management flow", level=1)
for text in [
    "Open Admin Banners at /admin/banners and choose Add Banner or edit an existing banner.",
    "Choose a Placement. Available placements are Home slider, Home Offer Zone, Shop slider, Search results slider and one slider placement for every active category.",
    "Upload a JPG, PNG or WebP image up to 2 MB. The editor recommends 1600 by 600 pixels for sliders and 1000 by 400 pixels for Offer Zone cards.",
    "Set Position, destination link and Enabled status. Lower Position values appear first.",
    "The Banner database table stores image, title, subtitle, link, placement, active state and sort order. Storefront queries only active records.",
]:
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(text)

doc.add_heading("Admin managed banner sources", level=1)
add_table(
    ["Storefront area", "Placement value", "Source and priority", "Fallback"],
    [
        ("Home hero slider", "home", "Banner table through getVayziqHomeData. Active banners sorted by sortOrder then createdAt. Up to 8 are loaded.", "Four images under public/vayziq including hero paired, women, men and fashion grid."),
        ("Home editorial offer strips", "offer-zone", "The first two active Offer Zone banner records are used by VayziqHome as the two existing offer strips.", "Women and men category images with built in offer copy."),
        ("Home coupon Offer Zone", "Not a Banner placement", "Coupon table through getVayziqHomeData. Only active, currently valid and usage available coupons are shown.", "No coupon cards if there are no eligible coupons."),
        ("Shop page banner", "shop", "CollectionPage calls getActiveBanner and displays the first active banner.", "public/vayziq/hero-paired-v2.png with the Move freely message."),
        ("Search results banner", "search", "The Search results placement is supported by the Banner system and editor.", "Page specific search presentation when no managed banner is supplied."),
        ("Category page banner", "category:{slug}", "Each active category gets its own placement, for example category:men. CollectionPage loads the first active banner.", "Code selects women, jogger, T shirt, fashion grid or men imagery from public/vayziq based on category name and slug."),
        ("Offer editorial cards", "offer-zone", "The OfferZone component calls getActiveBanners for all active Offer Zone records.", "Two demo offer cards use images under public/uploads/banners."),
    ],
    [1.35, 1.05, 2.55, 2.05],
)

doc.add_heading("Code based page banners", level=1)
add_table(
    ["Page or feature", "Component", "Current source", "How to change"],
    [
        ("About page", "VayziqPageBanner", "public/vayziq/hero-paired-v2.png passed directly by the page.", "Change the image prop in app/(store)/about/page.tsx or make the component database driven."),
        ("Contact page", "VayziqPageBanner", "public/vayziq/category-women.png passed directly by the page.", "Change the image prop in app/(store)/contact/page.tsx."),
        ("Legal pages", "LegalPageLayout and VayziqPageBanner", "Default public/vayziq/hero-paired-v2.png from VayziqPageBanner.", "Change the default image in VayziqPageBanner or add page specific image props."),
        ("Offers listing page", "OffersPage", "The compact black and yellow hero is built with CSS and text in app/(store)/offers/page.tsx. It does not use the Banner table.", "Edit the hero markup and classes in the offers page, or add a new Admin placement."),
        ("Individual offer page", "PageBanner", "Coupon based promo slide or an active offer-zone Banner record selected by numeric banner id.", "Edit promo slide generation or the linked Offer Zone banner in Admin."),
        ("BannerSlider image default", "BannerSlider", "public/vayziq/hero-paired-v2.png is used when a slide has no image.", "Supply a slide image or change the default in BannerSlider."),
        ("Watch and Buy posters", "WatchAndBuy", "Reel poster uploaded or selected in Admin Reels. Historical editorial examples reference public/uploads/banners/home-saree-editorial.png.", "Manage through Admin Reels; these are reel posters, not Banner records."),
    ],
    [1.35, 1.45, 2.5, 1.7],
)

doc.add_heading("Important distinction", level=1)
p = doc.add_paragraph()
p.add_run("Offer banners and coupon cards are different data. ").bold = True
p.add_run("The two image based offer strips use Banner records with placement offer-zone. The new coupon style Offers Zone uses Coupon records and derives its heading, condition, description and code from the coupon configuration. Removing or changing a Banner does not remove a Coupon, and changing a Coupon does not replace the editorial banner images.")

doc.add_heading("Source files", level=1)
sources = [
    ("lib/vayziq-home-data.ts", "Home hero banners, first two Offer Zone banners, coupon cards and fallbacks."),
    ("components/home/VayziqHome.tsx", "Home placement and rendering of hero, editorial offer strips and coupon Offer Zone."),
    ("components/home/OfferZone.module.css", "Visual design for the coupon Offer Zone."),
    ("lib/storefront.ts", "Banner placement types and active banner database queries."),
    ("components/admin/BannerEditor.tsx", "Admin placement choices, upload instructions, destination and activation controls."),
    ("components/store/CollectionPage.tsx", "Shop and category banner selection plus code fallbacks."),
    ("components/store/BannerSlider.tsx", "Shared banner rendering and default image behavior."),
    ("components/store/VayziqPageBanner.tsx", "About, Contact and legal page banner layout and default image."),
    ("app/(store)/offers/page.tsx", "Code based Offers page hero and live offer card layout."),
    ("prisma/schema.prisma", "Banner and Coupon database models."),
]
for path, desc in sources:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(path)
    r.bold = True
    p.add_run(f"  {desc}")

doc.add_heading("Recommended future cleanup", level=1)
for text in [
    "Add dedicated Admin placements for About, Contact, Legal and Offers page heroes if non developers need to change them.",
    "Keep editorial Offer Zone banners and coupon cards as separate controls because they serve different purposes.",
    "Use one documented image ratio per placement and preview desktop and mobile crops before enabling a banner.",
    "Rename the Admin label Home Offer Zone to Offer Editorial Banners if users confuse it with coupon cards.",
]:
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(text)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Vayziq storefront banner architecture")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(110, 110, 110)

doc.core_properties.title = "Vayziq Banner Source Guide"
doc.core_properties.subject = "Banner data sources and management paths"
doc.core_properties.author = "Vayziq"
doc.save(OUT)
print(OUT)
